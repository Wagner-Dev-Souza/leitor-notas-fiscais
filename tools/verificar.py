#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Verificador de idempotencia do pipeline - frente F3 (preguica).

Roda o comando unico congelado (secao 2 do contrato) DUAS vezes e prova a
idempotencia comparando a contagem de linhas da planilha antes e depois:

    .venv/Scripts/python.exe tools/verificar.py

Regra do aceite (secao 9.3 do contrato): rodar de novo tem de dar ZERO linha
duplicada. Se a contagem mudar entre as rodadas, ou se aparecer `pedido_id`
repetido, o script sai com codigo != 0 - e o PO/QA usam isso como gate.

Codigos de saida:
    0  idempotencia provada (contagem identica, nenhuma duplicata)
    1  FALHOU: a contagem mudou entre as rodadas (duplicou linha)
    2  o comando do pipeline falhou (nao deu para medir)
    3  o pipeline rodou mas nao produziu os artefatos de saida

Escopo: este script nao escreve em `app/**`. Ele e apenas o instrumento de
medicao do PO e do QA.
"""

from __future__ import annotations

import argparse
import csv
import json
import shlex
import shutil
import subprocess
import sys
import time
from pathlib import Path
from typing import Optional

RAIZ = Path(__file__).resolve().parents[1]
if str(RAIZ) not in sys.path:
    sys.path.insert(0, str(RAIZ))

from app.contratos import COLUNAS_PLANILHA  # noqa: E402

XLSX = "controle_financeiro.xlsx"
CSV = "controle_financeiro.csv"
AUDITORIA = "auditoria.jsonl"
EXCECOES = "fila_excecoes.json"

COL_PEDIDO = COLUNAS_PLANILHA.index("pedido_id")
COL_STATUS = COLUNAS_PLANILHA.index("status_validacao")


# ------------------------------------------------------------------ medicao


def contar_xlsx(caminho: Path) -> dict:
    from openpyxl import load_workbook

    if not caminho.exists():
        return {"existe": False}
    wb = load_workbook(caminho, read_only=True, data_only=True)
    ws = wb.active
    linhas = []
    cabecalho = None
    for i, linha in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            cabecalho = [str(c) if c is not None else "" for c in linha]
            continue
        if linha is None or all(c is None or str(c).strip() == "" for c in linha):
            continue
        linhas.append(linha)
    wb.close()

    pedidos = [str(l[COL_PEDIDO]) for l in linhas if len(l) > COL_PEDIDO and l[COL_PEDIDO]]
    status = {}
    for l in linhas:
        if len(l) > COL_STATUS and l[COL_STATUS]:
            status[str(l[COL_STATUS])] = status.get(str(l[COL_STATUS]), 0) + 1
    repetidos = sorted({p for p in pedidos if pedidos.count(p) > 1})
    return {
        "existe": True,
        "linhas": len(linhas),
        "pedidos": len(pedidos),
        "pedidos_distintos": len(set(pedidos)),
        "pedidos_repetidos": repetidos,
        "cabecalho": cabecalho,
        "status": status,
    }


def contar_csv(caminho: Path) -> dict:
    if not caminho.exists():
        return {"existe": False}
    with open(caminho, newline="", encoding="utf-8") as fh:
        linhas = [l for l in csv.reader(fh) if l and any(c.strip() for c in l)]
    return {"existe": True, "linhas": max(0, len(linhas) - 1)}   # menos o cabecalho


def contar_jsonl(caminho: Path) -> int:
    if not caminho.exists():
        return 0
    with open(caminho, encoding="utf-8") as fh:
        return sum(1 for l in fh if l.strip())


def contar_excecoes(caminho: Path) -> int:
    if not caminho.exists():
        return 0
    dado = json.loads(caminho.read_text(encoding="utf-8"))
    if isinstance(dado, list):
        return len(dado)
    if isinstance(dado, dict):
        itens = dado.get("itens", dado.get("excecoes", dado.get("pendencias")))
        if isinstance(itens, list):
            return len(itens)
        return len(dado)
    return 0


def medir(out: Path) -> dict:
    xlsx = contar_xlsx(out / XLSX)
    return {
        "xlsx": xlsx,
        "csv": contar_csv(out / CSV),
        "auditoria": contar_jsonl(out / AUDITORIA),
        "excecoes": contar_excecoes(out / EXCECOES),
    }


# -------------------------------------------------------------------- runners


def python_do_projeto() -> Optional[Path]:
    for rel in ("Scripts/python.exe", "bin/python"):
        p = RAIZ / ".venv" / rel
        if p.exists():
            return p
    return None


def montar_comando(args) -> list[str]:
    if args.comando:
        return shlex.split(args.comando, posix=True)
    py = Path(args.python) if args.python else python_do_projeto()
    if py is None:
        raise SystemExit("ERRO: nao achei o python do venv; passe --python ou --comando")
    return [str(py), "-m", "app.run", "--mock"]


def rodar(comando: list[str], rodada: int) -> tuple[int, float, str]:
    inicio = time.perf_counter()
    proc = subprocess.run(
        comando, cwd=str(RAIZ), capture_output=True, text=True,
        encoding="utf-8", errors="replace",
    )
    return proc.returncode, time.perf_counter() - inicio, (proc.stdout or "") + (proc.stderr or "")


# ---------------------------------------------------------------------- saida


def imprimir_rodada(rotulo: str, med: dict, segundos: float, codigo: int) -> None:
    x = med["xlsx"]
    if not x.get("existe"):
        print(f"{rotulo}: exit={codigo}  {segundos:.1f}s  |  {XLSX} NAO EXISTE")
        return
    print(
        f"{rotulo}: exit={codigo}  {segundos:.1f}s"
        f"  |  xlsx linhas={x['linhas']} (pedido_id distintos={x['pedidos_distintos']})"
        f"  |  csv linhas={med['csv'].get('linhas', '-')}"
        f"  |  auditoria={med['auditoria']}  excecoes={med['excecoes']}"
    )
    if x.get("status"):
        print(f"          status_validacao: {x['status']}")


def conferir_cabecalho(xlsx: dict) -> Optional[str]:
    cabecalho = xlsx.get("cabecalho")
    if not cabecalho:
        return None
    if tuple(cabecalho) != tuple(COLUNAS_PLANILHA):
        return (f"cabecalho da planilha DIVERGE de contratos.COLUNAS_PLANILHA\n"
                f"    esperado: {list(COLUNAS_PLANILHA)}\n"
                f"    obtido  : {cabecalho}")
    return None


def main(argv: Optional[list[str]] = None) -> int:
    ap = argparse.ArgumentParser(
        description="Roda o comando unico 2x e prova a idempotencia da planilha."
    )
    ap.add_argument("--comando", default=None,
                    help="comando completo (default: '<python-do-venv> -m app.run --mock')")
    ap.add_argument("--python", default=None, help="interpretador do projeto")
    ap.add_argument("--out", default=str(RAIZ / "data" / "out"), help="diretorio de saida")
    ap.add_argument("--inbox", default=str(RAIZ / "data" / "mocks"), help="diretorio de entrada")
    ap.add_argument("--rodadas", type=int, default=2, help="quantidade de rodadas (default: 2)")
    ap.add_argument("--limpar-out", action="store_true",
                    help="apaga data/out antes da primeira rodada (teste a frio)")
    ap.add_argument("--json", action="store_true", help="imprime o resumo em JSON no final")
    args = ap.parse_args(argv)

    out = Path(args.out)
    comando = montar_comando(args)

    print("=" * 78)
    print("VERIFICADOR DE IDEMPOTENCIA - tools/verificar.py (F3 / preguica)")
    print("=" * 78)
    print(f"comando : {' '.join(comando)}")
    print(f"cwd     : {RAIZ}")
    print(f"inbox   : {args.inbox}")
    print(f"out     : {out}")
    print(f"rodadas : {args.rodadas}")
    print("-" * 78)

    if args.limpar_out and out.exists():
        shutil.rmtree(out)
        print(f"[--limpar-out] removido: {out}")

    resultado: dict = {"comando": comando, "rodadas": [], "ok": False}
    medicoes: list[dict] = []

    for i in range(1, args.rodadas + 1):
        codigo, segundos, saida = rodar(comando, i)
        med = medir(out)
        medicoes.append(med)
        imprimir_rodada(f"rodada {i}", med, segundos, codigo)
        resultado["rodadas"].append({"rodada": i, "exit": codigo, "segundos": round(segundos, 2),
                                     **med})
        if codigo != 0:
            print("-" * 78)
            print(f"FALHOU: o comando do pipeline saiu com codigo {codigo} na rodada {i}.")
            print("ultimas linhas da saida do pipeline:")
            for linha in saida.strip().splitlines()[-25:]:
                print(f"    | {linha}")
            resultado["ok"] = False
            resultado["motivo"] = f"comando_falhou_rodada_{i}"
            if args.json:
                print(json.dumps(resultado, ensure_ascii=False, indent=2, default=str))
            return 2

    primeira = medicoes[0]
    if not primeira["xlsx"].get("existe"):
        print("-" * 78)
        print(f"FALHOU: o pipeline rodou mas {out / XLSX} nao existe.")
        resultado["motivo"] = "sem_planilha"
        if args.json:
            print(json.dumps(resultado, ensure_ascii=False, indent=2, default=str))
        return 3

    print("-" * 78)
    base = primeira["xlsx"]["linhas"]
    problemas: list[str] = []

    for i, med in enumerate(medicoes[1:], start=2):
        atual = med["xlsx"].get("linhas")
        if atual != base:
            problemas.append(
                f"contagem da planilha MUDOU: rodada 1 = {base} linhas, rodada {i} = {atual}"
            )
        if med["csv"].get("linhas") not in (None, primeira["csv"].get("linhas")):
            problemas.append(
                f"contagem do CSV MUDOU: rodada 1 = {primeira['csv'].get('linhas')}, "
                f"rodada {i} = {med['csv'].get('linhas')}"
            )

    repetidos = set()
    for med in medicoes:
        repetidos.update(med["xlsx"].get("pedidos_repetidos") or [])
    if repetidos:
        problemas.append(f"pedido_id repetido na planilha: {sorted(repetidos)}")

    aviso = conferir_cabecalho(primeira["xlsx"])
    if aviso:
        print(f"AVISO: {aviso}")

    pedidos_distintos = primeira["xlsx"]["pedidos_distintos"]
    if pedidos_distintos != base:
        print(f"AVISO: {base} linhas para {pedidos_distintos} pedido_id distintos.")

    if problemas:
        print("RESULTADO: FALHOU (linha duplicada / contagem instavel)")
        for p in problemas:
            print(f"  - {p}")
        resultado["ok"] = False
        resultado["problemas"] = problemas
        if args.json:
            print(json.dumps(resultado, ensure_ascii=False, indent=2, default=str))
        return 1

    print(
        f"RESULTADO: PASSOU - {base} linha(s) na planilha em TODAS as {args.rodadas} rodadas, "
        f"{pedidos_distintos} pedido_id distintos, zero duplicata."
    )
    print(f"  xlsx: {out / XLSX}")
    print(f"  csv : {out / CSV}")
    print(f"  trilha: {out / AUDITORIA} ({primeira['auditoria']} registros) | "
          f"excecoes: {out / EXCECOES} ({primeira['excecoes']})")
    resultado["ok"] = True
    resultado["linhas"] = base
    if args.json:
        print(json.dumps(resultado, ensure_ascii=False, indent=2, default=str))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

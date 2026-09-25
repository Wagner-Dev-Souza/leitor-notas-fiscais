"""Aprovacao na planilha: modulo DORMENTE desde 25/09/2026.

O PO decidiu lancamento direto (toda nota entra na planilha com destaque visual), entao o pipeline
nao chama mais `app/aprovacao.py` e a rodada NAO cria aba de revisao. Este arquivo guarda duas
coisas: a prova de que a aba nao volta sozinha, e a cobertura do modulo para o caso de o fluxo de
aprovacao ser religado.
"""

from __future__ import annotations

import json
import os
import sqlite3
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conftest import MOCKS, RAIZ_PROJETO

RAIZ = Path(__file__).resolve().parent.parent
if str(RAIZ) not in sys.path:
    sys.path.insert(0, str(RAIZ))

from app import aprovacao  # noqa: E402

PYTHON = str(RAIZ / ".venv" / "Scripts" / "python.exe")


def rodar_mock(out: Path) -> subprocess.CompletedProcess:
    return subprocess.run(
        [
            PYTHON,
            "-m",
            "app.run",
            "--mock",
            "--inbox",
            str(MOCKS),
            "--out",
            str(out),
            "--db",
            str(out / "pipeline.db"),
        ],
        cwd=str(RAIZ_PROJETO),
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        env={**os.environ, "PYTHONPATH": str(RAIZ_PROJETO)},
        timeout=300,
    )


@pytest.fixture()
def rodada(tmp_path):
    out = tmp_path / "out"
    out.mkdir()
    resultado = rodar_mock(out)
    assert resultado.returncode == 0, resultado.stdout + resultado.stderr
    return out


def abrir_banco(out: Path) -> sqlite3.Connection:
    conn = sqlite3.connect(out / "pipeline.db")
    conn.row_factory = sqlite3.Row
    return conn


# ------------------------------------------------------------------ o fluxo antigo esta desligado


def test_rodada_nao_cria_aba_de_revisao(rodada):
    """Lancamento direto: a planilha tem UMA aba - a oficial."""
    caminho = rodada / "controle_financeiro.xlsx"
    wb = load_workbook(caminho)
    assert wb.sheetnames == [aprovacao.ABA_OFICIAL]
    assert wb.active.title == aprovacao.ABA_OFICIAL


def test_rodada_apaga_aba_de_revisao_de_uma_planilha_antiga(rodada):
    """Sobrou a aba de uma rodada antiga? A rodada seguinte limpa - duas verdades confundem."""
    from openpyxl import load_workbook as carregar

    caminho = rodada / "controle_financeiro.xlsx"
    wb = carregar(caminho)
    wb.create_sheet(aprovacao.ABA_REVISAO)
    wb.save(caminho)

    assert rodar_mock(rodada).returncode == 0
    assert carregar(caminho).sheetnames == [aprovacao.ABA_OFICIAL]


# ------------------------------------------------------------------ o modulo em si


def test_rotulo_de_decisao_aceita_o_que_o_humano_digita():
    for bruto in ("APROVAR", "aprovar", "Aprovado", "ok", "SIM", "s", "X"):
        assert aprovacao.rotulo_decisao(bruto) == aprovacao.DECISAO_APROVAR, bruto
    for bruto in ("REJEITAR", "rejeitado", "nao", "NÃO", "N"):
        assert aprovacao.rotulo_decisao(bruto) == aprovacao.DECISAO_REJEITAR, bruto
    for bruto in (None, "", "   ", "talvez", 42):
        assert aprovacao.rotulo_decisao(bruto) is None, bruto


def _montar_aba(rodada: Path) -> Path:
    caminho = rodada / "controle_financeiro.xlsx"
    conn = abrir_banco(rodada)
    try:
        linhas = aprovacao.exportar_revisao(conn, caminho)
    finally:
        conn.close()
    assert linhas >= 1, "a rodada mock tem pendencias para revisar"
    return caminho


def _decidir(caminho: Path, decisao: str, pedido_id: str | None = None, **campos) -> str:
    wb = load_workbook(caminho)
    ws = wb[aprovacao.ABA_REVISAO]
    cabecalho = [celula.value for celula in ws[1]]
    for linha in ws.iter_rows(min_row=3):
        valor = linha[cabecalho.index("pedido_id")].value
        if pedido_id is None or valor == pedido_id:
            linha[cabecalho.index("DECISAO")].value = decisao
            linha[cabecalho.index("REVISOR")].value = "Wagner"
            for nome, conteudo in campos.items():
                linha[cabecalho.index(nome)].value = conteudo
            wb.save(caminho)
            return str(valor)
    raise AssertionError("nenhuma linha da aba Revisao para decidir")


def _linha_oficial(caminho: Path, pedido_id: str) -> dict | None:
    ws = load_workbook(caminho)[aprovacao.ABA_OFICIAL]
    cabecalho = [celula.value for celula in ws[1]]
    for linha in ws.iter_rows(min_row=2):
        valores = dict(zip(cabecalho, [celula.value for celula in linha]))
        if valores["pedido_id"] == pedido_id:
            return valores
    return None


def test_aprovar_promove_o_pedido_e_fecha_a_pendencia(rodada):
    caminho = _montar_aba(rodada)
    pedido_id = _decidir(caminho, "APROVAR", VALOR_TOTAL_CORRIGIDO="1.234,56")

    conn = abrir_banco(rodada)
    try:
        resumo = aprovacao.aplicar_decisoes(conn, caminho)
        pedido = conn.execute("SELECT * FROM pedidos WHERE id = ?", (pedido_id,)).fetchone()
        pendencia = conn.execute(
            "SELECT * FROM fila_excecoes WHERE pedido_id = ?", (pedido_id,)
        ).fetchone()
    finally:
        conn.close()

    assert resumo["aprovadas"] == 1
    assert resumo["pendencias_fechadas"] >= 1
    assert pedido["status"] == "validado"
    assert pedido["valor_total_centavos"] == 123456, "a correcao digitada nao foi usada"
    assert pendencia["status"] == "resolvida"
    assert pendencia["resolvida_por"] == "Wagner"
    assert pendencia["decisao"] == "aprovado"


def test_rejeitar_fecha_a_pendencia_com_a_decisao(rodada):
    caminho = _montar_aba(rodada)
    pedido_id = _decidir(caminho, "REJEITAR")

    conn = abrir_banco(rodada)
    try:
        resumo = aprovacao.aplicar_decisoes(conn, caminho)
        pedido = conn.execute("SELECT status FROM pedidos WHERE id = ?", (pedido_id,)).fetchone()
        pendencia = conn.execute(
            "SELECT * FROM fila_excecoes WHERE pedido_id = ?", (pedido_id,)
        ).fetchone()
    finally:
        conn.close()

    assert resumo["rejeitadas"] == 1
    assert pedido["status"] == "rejeitado"
    assert pendencia["status"] == "resolvida" and pendencia["decisao"] == "rejeitado"


def test_aprovacao_sem_valor_nao_vira_linha(rodada):
    """Nenhum valor (lido ou digitado): aprovar nao inventa numero - fica aviso e a fila aberta."""
    caminho = _montar_aba(rodada)
    conn = abrir_banco(rodada)
    try:
        sem_valor = conn.execute(
            """
            SELECT f.pedido_id FROM fila_excecoes f JOIN pedidos p ON p.id = f.pedido_id
             WHERE p.valor_total_centavos IS NULL AND f.status = 'aberta' LIMIT 1
            """
        ).fetchone()
    finally:
        conn.close()
    assert sem_valor is not None, "o corpus mock tem de ter pendencia sem valor lido"
    pedido_id = _decidir(caminho, "APROVAR", pedido_id=str(sem_valor["pedido_id"]))

    conn = abrir_banco(rodada)
    try:
        resumo = aprovacao.aplicar_decisoes(conn, caminho)
        pedido = conn.execute("SELECT status FROM pedidos WHERE id = ?", (pedido_id,)).fetchone()
        pendencia = conn.execute(
            "SELECT status FROM fila_excecoes WHERE pedido_id = ?", (pedido_id,)
        ).fetchone()
    finally:
        conn.close()

    assert resumo["aprovadas"] == 0
    assert any("sem valor total" in aviso for aviso in resumo["avisos"]), resumo["avisos"]
    assert pedido["status"] != "validado"
    assert pendencia["status"] == "aberta"


def test_decidir_pedido_inexistente_vira_aviso_sem_quebrar(rodada):
    caminho = _montar_aba(rodada)
    _decidir(caminho, "APROVAR")
    wb = load_workbook(caminho)
    ws = wb[aprovacao.ABA_REVISAO]
    cabecalho = [celula.value for celula in ws[1]]
    ws.cell(row=3, column=cabecalho.index("pedido_id") + 1, value="ped_inventado")
    wb.save(caminho)

    decisoes = aprovacao.ler_decisoes(caminho)
    assert decisoes, "a decisao digitada tem de ser lida"
    conn = abrir_banco(rodada)
    try:
        resumo = aprovacao.aplicar_decisoes(conn, caminho)
    finally:
        conn.close()
    assert resumo["ignoradas"] >= 1
    assert any("nao existe no banco" in aviso for aviso in resumo["avisos"])


def test_csv_e_planilha_seguem_com_a_mesma_contagem(rodada):
    """Sanidade da evidencia: o destaque nao pode quebrar a contagem linha a linha."""
    resumo = json.loads((rodada / "resumo.json").read_text(encoding="utf-8"))
    ws = load_workbook(rodada / "controle_financeiro.xlsx")[aprovacao.ABA_OFICIAL]
    csv_bruto = (rodada / "controle_financeiro.csv").read_text(encoding="utf-8").splitlines()
    assert ws.max_row - 1 == resumo["linhas_planilha"] == len(csv_bruto) - 1

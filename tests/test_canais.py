"""Coleta dos canais reais (`app/canais.py`, F8) e o ciclo do modo real.

Contrato: `docs/execucao/00b-contrato-fechamento.md` secoes 2 (D3, D4, D5) e 3.2/3.3.
Regra que vale para todo este arquivo: **nenhuma chamada a `api.telegram.org` ou
`graph.facebook.com`**. O que se prova aqui e a montagem da requisicao, a gravacao do
envelope no MESMO formato dos mocks e o ciclo completo contra stub local `127.0.0.1` -
que e exatamente o que a decisao D4 autoriza (nao existe token nem numero real nesta
entrega, e nada disso e apresentado como chamada credencial real).

O seam `TransporteHTTP` e injetado por parametro (sem `monkeypatch`); o caminho real
`transporte_urllib()` tambem e exercitado de verdade contra o stub HTTP local.
"""

from __future__ import annotations

import json
import os
import shutil
import socket
import subprocess
import sys
import threading
from contextlib import contextmanager
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import pytest

from app import canais
from app import config as CFG
from conftest import MOCKS, RAIZ_PROJETO

TELEGRAM_TOKEN = "1111111111:TOKEN_FICTICIO_DE_TESTE"
TELEGRAM_CHAT_ID = "-1001234567890"
WHATSAPP_TOKEN = "EAAG_TOKEN_FICTICIO_DE_TESTE"
WHATSAPP_VERIFY = "VERIFY_FICTICIO_DE_TESTE"

PYTHON = str(Path(sys.executable))
if not (RAIZ_PROJETO / ".venv" / "Scripts" / "python.exe").is_file():  # pragma: no cover
    PYTHON = str(Path(sys.executable))


# --------------------------------------------------------------------- utilidades


def env_de_teste(tmp_path: Path, **campos) -> Path:
    """`.env` de teste com TODOS os caminhos dentro de `tmp_path` (nada toca o repositorio)."""
    valores = {
        "MODO_EXECUCAO": "real",
        "CANAIS_ATIVOS": "telegram",
        "TELEGRAM_BOT_TOKEN": TELEGRAM_TOKEN,
        "TELEGRAM_CHAT_ID": TELEGRAM_CHAT_ID,
        "TELEGRAM_TIMEOUT_S": "5",
        "WHATSAPP_TOKEN": WHATSAPP_TOKEN,
        "WHATSAPP_PHONE_NUMBER_ID": "1234567890",
        "WHATSAPP_VERIFY_TOKEN": WHATSAPP_VERIFY,
        "INBOX_DIR": str(tmp_path / "inbox"),
        "OUT_DIR": str(tmp_path / "out"),
        "LOG_DIR": str(tmp_path / "logs"),
        "WHATSAPP_WEBHOOK_DIR": str(tmp_path / "webhook"),
    }
    valores.update({k: v for k, v in campos.items() if v is not None})
    caminho = tmp_path / "env.teste"
    caminho.write_text("".join(f"{k}={v}\n" for k, v in valores.items()), encoding="utf-8")
    return caminho


def config_de_teste(tmp_path: Path, **campos) -> CFG.Config:
    return CFG.carregar(env_path=env_de_teste(tmp_path, **campos), ambiente={})


def update_telegram(update_id: int, chat_id=TELEGRAM_CHAT_ID, texto="Pedido 7501 confirmado, total R$ 480,00"):
    return {
        "update_id": update_id,
        "message": {
            "message_id": update_id + 100,
            "from": {"id": 55555, "is_bot": False, "first_name": "Maria", "username": "maria_compras"},
            "chat": {"id": int(chat_id), "title": "Compras Fornecedores", "type": "group"},
            "date": 1758200000,
            "text": texto,
        },
    }


def envelope_whatsapp(texto="Pedido 7502 fechado, total R$ 310,00", id_externo="wamid.COLETADO"):
    return {
        "object": "whatsapp_business_account",
        "entry": [
            {
                "id": "WABA-TESTE",
                "changes": [
                    {
                        "field": "messages",
                        "value": {
                            "messaging_product": "whatsapp",
                            "metadata": {"display_phone_number": "551140028922", "phone_number_id": "PH-TESTE"},
                            "contacts": [{"profile": {"name": "Jose da Silva"}, "wa_id": "5511998887777"}],
                            "messages": [
                                {
                                    "from": "5511998887777",
                                    "id": id_externo,
                                    "timestamp": "1758200000",
                                    "type": "text",
                                    "text": {"body": texto},
                                }
                            ],
                        },
                    }
                ],
            }
        ],
    }


class TransporteFalso:
    """Seam `TransporteHTTP` injetado pelo teste: registra a chamada e devolve o roteiro."""

    def __init__(self, resposta=None, erro=None):
        self.resposta = resposta if resposta is not None else {"ok": True, "result": []}
        self.erro = erro
        self.chamadas: list[dict] = []

    def get_json(self, url, params=None, cabecalhos=None, timeout=30):
        self.chamadas.append({"metodo": "GET", "url": url, "params": params, "timeout": timeout})
        if self.erro is not None:
            raise self.erro
        return self.resposta

    def post_json(self, url, corpo, cabecalhos=None, timeout=30):  # pragma: no cover - seam simetrico
        self.chamadas.append({"metodo": "POST", "url": url, "corpo": corpo, "timeout": timeout})
        if self.erro is not None:
            raise self.erro
        return self.resposta


def porta_livre() -> int:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return int(s.getsockname()[1])


@contextmanager
def stub_http(respondedor):
    """Servidor HTTP local. `respondedor(caminho, consulta, corpo)` -> (status, corpo)."""
    recebidas: list[dict] = []

    class Handler(BaseHTTPRequestHandler):
        def log_message(self, *args):  # silencia o ruido do http.server
            return

        def _responder(self, metodo: str, corpo_bruto: bytes = b""):
            partes = urlparse(self.path)
            recebidas.append({"metodo": metodo, "caminho": partes.path, "consulta": parse_qs(partes.query)})
            status, corpo = respondedor(partes.path, parse_qs(partes.query), corpo_bruto)
            dados = corpo.encode("utf-8") if isinstance(corpo, str) else json.dumps(corpo, ensure_ascii=False).encode("utf-8")
            self.send_response(status)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Content-Length", str(len(dados)))
            self.end_headers()
            self.wfile.write(dados)

        def do_GET(self):  # noqa: N802
            self._responder("GET")

        def do_POST(self):  # noqa: N802
            tamanho = int(self.headers.get("Content-Length") or 0)
            self._responder("POST", self.rfile.read(tamanho))

    servidor = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    porta = servidor.server_address[1]
    thread = threading.Thread(target=servidor.serve_forever, daemon=True)
    thread.start()
    try:
        yield porta, recebidas
    finally:
        servidor.shutdown()
        servidor.server_close()
        thread.join(timeout=5)


def esperar_no_ar(porta: int, prazo_s: float = 15.0) -> None:
    """Curto poll de conectividade (sem sleep longo): o servidor sobe em milissegundos."""
    from time import monotonic, sleep

    fim = monotonic() + prazo_s
    while monotonic() < fim:
        try:
            with socket.create_connection(("127.0.0.1", porta), timeout=0.25):
                return
        except OSError:
            sleep(0.05)
    raise AssertionError(f"servidor local nao subiu na porta {porta} em {prazo_s}s")


def ler_jsonl(caminho: Path) -> list[dict]:
    return [json.loads(linha) for linha in caminho.read_text(encoding="utf-8").splitlines() if linha.strip()]


# --------------------------------------------------------------------- telegram: envelopes


def test_coletar_telegram_grava_um_jsonl_por_coleta_no_formato_do_mock(tmp_path):
    cfg = config_de_teste(tmp_path)
    transporte = TransporteFalso(
        {
            "ok": True,
            "result": [update_telegram(101), update_telegram(102, texto="Pedido 7503 confirmado, total R$ 99,90")],
        }
    )
    resultado = canais.coletar_telegram(cfg, transporte=transporte)

    assert resultado.canal == "telegram"
    assert resultado.mensagens == 2
    assert resultado.gravou is True
    assert len(resultado.arquivos) == 1
    arquivo = resultado.arquivos[0]
    assert arquivo.parent == Path(cfg.inbox_dir) / "telegram"
    assert arquivo.name.startswith("telegram_coleta_") and arquivo.suffix == ".jsonl"

    linhas = ler_jsonl(arquivo)
    assert len(linhas) == 2, "o numero de linhas tem de bater com o numero de updates aceitos"
    assert linhas[0]["update_id"] == 101

    # mesmo formato do mock: mesmas chaves do update real de data/mocks
    mock = ler_jsonl(MOCKS / "telegram" / "telegram_1.jsonl")[0]
    assert set(linhas[0]) == set(mock), "o envelope coletado saiu com formato diferente do mock"
    assert set(linhas[0]["message"]) >= set(mock["message"]) - {"message_id"}
    assert linhas[0]["message"]["text"] == "Pedido 7501 confirmado, total R$ 480,00"


def test_envelope_gravado_preserva_acentuacao(tmp_path):
    cfg = config_de_teste(tmp_path)
    transporte = TransporteFalso({"ok": True, "result": [update_telegram(1, texto="Pedido 7504 - entrega em São Paulo, 2ª via")]})
    resultado = canais.coletar_telegram(cfg, transporte=transporte)
    bruto = resultado.arquivos[0].read_text(encoding="utf-8")
    assert "São Paulo" in bruto and "2ª" in bruto, "ensure_ascii=False nao foi respeitado"
    assert "\\u00e3" not in bruto


def test_update_sem_message_nao_vira_lixo_no_arquivo(tmp_path):
    cfg = config_de_teste(tmp_path)
    transporte = TransporteFalso(
        {
            "ok": True,
            "result": [
                {"update_id": 1},
                {"update_id": 2, "message": None},
                "lixo",
                None,
                {"update_id": 3, "message": {"chat": {"id": int(TELEGRAM_CHAT_ID)}, "text": "Pedido 7505 total R$ 10,00"}},
            ],
        }
    )
    resultado = canais.coletar_telegram(cfg, transporte=transporte)
    assert resultado.mensagens == 1
    linhas = ler_jsonl(resultado.arquivos[0])
    assert [linha["update_id"] for linha in linhas] == [3]


def test_filtro_de_chat_id_descarta_conversa_nao_autorizada(tmp_path):
    cfg = config_de_teste(tmp_path)
    transporte = TransporteFalso(
        {
            "ok": True,
            "result": [
                update_telegram(1, chat_id="-999999999"),
                update_telegram(2, chat_id=TELEGRAM_CHAT_ID),
            ],
        }
    )
    resultado = canais.coletar_telegram(cfg, transporte=transporte)
    assert resultado.mensagens == 1
    assert ler_jsonl(resultado.arquivos[0])[0]["update_id"] == 2


def test_sem_update_do_chat_configurado_nao_grava_e_nao_e_erro(tmp_path):
    cfg = config_de_teste(tmp_path)
    transporte = TransporteFalso({"ok": True, "result": [update_telegram(1, chat_id="-999999999")]})
    resultado = canais.coletar_telegram(cfg, transporte=transporte)
    assert resultado.arquivos == ()
    assert resultado.mensagens == 0
    assert "chat configurado" in resultado.detalhe
    assert list((Path(cfg.inbox_dir) / "telegram").glob("*.jsonl")) == []


def test_sem_update_nenhum_nao_e_erro(tmp_path):
    cfg = config_de_teste(tmp_path)
    resultado = canais.coletar_telegram(cfg, transporte=TransporteFalso({"ok": True, "result": []}))
    assert resultado.arquivos == ()
    assert resultado.mensagens == 0
    assert "sem update" in resultado.detalhe and "nao e erro" in resultado.detalhe


def test_dois_updates_no_mesmo_segundo_nao_se_sobrescrevem(tmp_path):
    cfg = config_de_teste(tmp_path)
    primeiro = canais.coletar_telegram(cfg, transporte=TransporteFalso({"ok": True, "result": [update_telegram(1)]}))
    segundo = canais.coletar_telegram(cfg, transporte=TransporteFalso({"ok": True, "result": [update_telegram(2)]}))
    assert primeiro.arquivos[0] != segundo.arquivos[0]
    assert primeiro.arquivos[0].exists() and segundo.arquivos[0].exists()


# --------------------------------------------------------------------- telegram: requisicao


def test_montagem_da_requisicao_get_updates(tmp_path):
    cfg = config_de_teste(tmp_path, TELEGRAM_API_BASE="https://exemplo.invalido/api")
    transporte = TransporteFalso({"ok": True, "result": []})

    canais.coletar_telegram(cfg, transporte=transporte)
    chamada = transporte.chamadas[0]
    assert chamada["metodo"] == "GET"
    assert chamada["url"] == f"https://exemplo.invalido/api/bot{TELEGRAM_TOKEN}/getUpdates"
    assert chamada["params"] is None
    assert chamada["timeout"] == 5.0

    canais.coletar_telegram(cfg, transporte=transporte, offset=777)
    assert transporte.chamadas[1]["params"] == {"offset": 777}


def test_coleta_usa_urllib_real_contra_stub_local_127_0_0_1(tmp_path):
    """O caminho real (`transporte_urllib`) exercitado de verdade, sem sair da maquina."""
    with stub_http(lambda caminho, consulta, corpo: (200, {"ok": True, "result": [update_telegram(9)]})) as (porta, recebidas):
        cfg = config_de_teste(tmp_path, TELEGRAM_API_BASE=f"http://127.0.0.1:{porta}")
        resultado = canais.coletar_telegram(cfg)  # sem transporte injetado: urllib de verdade

    assert resultado.mensagens == 1
    assert ler_jsonl(resultado.arquivos[0])[0]["update_id"] == 9
    assert recebidas[0]["caminho"] == f"/bot{TELEGRAM_TOKEN}/getUpdates"


def test_http_401_vira_erro_canal_claro_e_sem_vazar_token(tmp_path):
    def responder(caminho, consulta, corpo):
        return 401, {"ok": False, "description": "Unauthorized"}

    with stub_http(responder) as (porta, _recebidas):
        cfg = config_de_teste(tmp_path, TELEGRAM_API_BASE=f"http://127.0.0.1:{porta}")
        with pytest.raises(canais.ErroCanal) as capturado:
            canais.coletar_telegram(cfg)

    mensagem = str(capturado.value)
    assert "telegram" in mensagem and "401" in mensagem
    assert "TELEGRAM_BOT_TOKEN" in mensagem, "a mensagem tem de dizer QUAL variavel conferir"
    assert TELEGRAM_TOKEN not in mensagem, "o token vazou na mensagem de erro"
    assert f"bot{TELEGRAM_TOKEN}" not in mensagem, "a URL com token vazou na mensagem de erro"


def test_falha_de_rede_vira_erro_canal_sem_url_credenciada(tmp_path):
    """Transporte real apontado para uma porta morta: falha local, mensagem clara."""
    cfg = config_de_teste(tmp_path, TELEGRAM_API_BASE="http://127.0.0.1:9", TELEGRAM_TIMEOUT_S="3")
    with pytest.raises(canais.ErroCanal) as capturado:
        canais.coletar_telegram(cfg)  # urllib de verdade, contra porta fechada

    mensagem = str(capturado.value)
    assert "telegram" in mensagem
    assert "nao alcancou" in mensagem
    assert TELEGRAM_TOKEN not in mensagem
    assert "no modo mock nada disso e necessario" in mensagem


def test_excecao_generica_do_transporte_vira_erro_canal_citando_as_variaveis(tmp_path):
    cfg = config_de_teste(tmp_path)
    transporte = TransporteFalso(erro=RuntimeError("falha inesperada do transporte"))
    with pytest.raises(canais.ErroCanal) as capturado:
        canais.coletar_telegram(cfg, transporte=transporte)
    mensagem = str(capturado.value)
    assert "falhou" in mensagem and "RuntimeError" in mensagem
    assert "TELEGRAM_BOT_TOKEN" in mensagem and "TELEGRAM_API_BASE" in mensagem
    assert TELEGRAM_TOKEN not in mensagem


def test_resposta_ok_false_vira_erro_canal(tmp_path):
    cfg = config_de_teste(tmp_path)
    transporte = TransporteFalso({"ok": False, "description": "Unauthorized"})
    with pytest.raises(canais.ErroCanal) as capturado:
        canais.coletar_telegram(cfg, transporte=transporte)
    assert "Unauthorized" in str(capturado.value)
    assert TELEGRAM_TOKEN not in str(capturado.value)


def test_coletar_telegram_sem_token_na_configuracao_e_erro_explicito(tmp_path):
    caminho = env_de_teste(tmp_path, TELEGRAM_BOT_TOKEN="", MODO_EXECUCAO="mock")
    cfg = CFG.carregar(env_path=caminho, ambiente={})
    with pytest.raises(canais.ErroCanal) as capturado:
        canais.coletar_telegram(cfg)
    assert "TELEGRAM_BOT_TOKEN" in str(capturado.value)
    assert "Traceback" not in str(capturado.value)


def test_detalhe_de_sucesso_nunca_contem_o_token(tmp_path):
    cfg = config_de_teste(tmp_path)
    resultado = canais.coletar_telegram(cfg, transporte=TransporteFalso({"ok": True, "result": [update_telegram(1)]}))
    assert TELEGRAM_TOKEN not in resultado.detalhe
    assert TELEGRAM_TOKEN not in resultado.arquivos[0].name
    assert resultado.arquivos[0].name in resultado.detalhe


# --------------------------------------------------------------------- whatsapp: webhook local


def test_ler_webhook_grava_envelope_com_entry_e_move_o_arquivo(tmp_path):
    cfg = config_de_teste(tmp_path)
    webhook_dir = Path(cfg.whatsapp.webhook_dir)
    webhook_dir.mkdir(parents=True, exist_ok=True)
    arquivo = webhook_dir / "20260919-101010_1.json"
    arquivo.write_text(json.dumps(envelope_whatsapp(), ensure_ascii=False), encoding="utf-8")

    resultado = canais.ler_webhook_whatsapp(cfg)

    assert resultado.canal == "whatsapp"
    assert resultado.mensagens == 1
    linhas = ler_jsonl(resultado.arquivos[0])
    assert linhas[0] == envelope_whatsapp()
    assert resultado.arquivos[0].parent == Path(cfg.inbox_dir) / "whatsapp"
    assert not arquivo.exists(), "com mover=True o arquivo consumido tem de sair da pasta quente"
    assert (webhook_dir / "processados" / arquivo.name).is_file()
    assert "processados" in resultado.detalhe


def test_envelope_de_status_sem_entry_e_descartado_com_explicacao(tmp_path):
    cfg = config_de_teste(tmp_path)
    webhook_dir = Path(cfg.whatsapp.webhook_dir)
    webhook_dir.mkdir(parents=True, exist_ok=True)
    (webhook_dir / "status.json").write_text(
        json.dumps({"object": "whatsapp_business_account", "entry": []}), encoding="utf-8"
    )

    resultado = canais.ler_webhook_whatsapp(cfg)

    assert resultado.arquivos == ()
    assert resultado.mensagens == 0
    assert "sem 'entry'" in resultado.detalhe
    assert (webhook_dir / "status.json").is_file(), "arquivo sem envelope nao deve ser descartado"
    assert not (webhook_dir / "processados").exists()


def test_mover_false_mantem_o_arquivo_na_pasta_do_webhook(tmp_path):
    cfg = config_de_teste(tmp_path)
    webhook_dir = Path(cfg.whatsapp.webhook_dir)
    webhook_dir.mkdir(parents=True, exist_ok=True)
    arquivo = webhook_dir / "msg.json"
    arquivo.write_text(json.dumps(envelope_whatsapp()), encoding="utf-8")

    resultado = canais.ler_webhook_whatsapp(cfg, mover=False)

    assert resultado.mensagens == 1
    assert arquivo.is_file(), "mover=False nao pode mover o arquivo"
    assert not (webhook_dir / "processados").exists()


def test_jsonl_com_varias_linhas_vira_varias_linhas_no_inbox(tmp_path):
    cfg = config_de_teste(tmp_path)
    webhook_dir = Path(cfg.whatsapp.webhook_dir)
    webhook_dir.mkdir(parents=True, exist_ok=True)
    conteudo = "\n".join(
        [
            json.dumps(envelope_whatsapp("Pedido 7601 total R$ 11,00", "wamid.A")),
            json.dumps({"object": "whatsapp_business_account", "entry": []}),
            json.dumps(envelope_whatsapp("Pedido 7602 total R$ 22,00", "wamid.B")),
        ]
    )
    (webhook_dir / "lote.jsonl").write_text(conteudo + "\n", encoding="utf-8")

    resultado = canais.ler_webhook_whatsapp(cfg)

    assert resultado.mensagens == 2
    assert len(ler_jsonl(resultado.arquivos[0])) == 2
    assert "1 sem 'entry' descartado(s)" in resultado.detalhe


def test_linha_invalida_nao_derruba_a_coleta(tmp_path):
    cfg = config_de_teste(tmp_path)
    webhook_dir = Path(cfg.whatsapp.webhook_dir)
    webhook_dir.mkdir(parents=True, exist_ok=True)
    (webhook_dir / "misto.jsonl").write_text(
        "{isso nao e json}\n" + json.dumps(envelope_whatsapp()) + "\n", encoding="utf-8"
    )

    resultado = canais.ler_webhook_whatsapp(cfg)

    assert resultado.mensagens == 1
    assert "1 linha(s)/arquivo(s) ilegiveis" in resultado.detalhe
    assert (webhook_dir / "processados" / "misto.jsonl").is_file()


def test_diretorio_de_webhook_ausente_nao_e_erro(tmp_path):
    cfg = config_de_teste(tmp_path, WHATSAPP_WEBHOOK_DIR=str(tmp_path / "nao_existe"))
    resultado = canais.ler_webhook_whatsapp(cfg)
    assert resultado.arquivos == ()
    assert "ainda nao existe" in resultado.detalhe
    assert "tools/receber_webhook_whatsapp.py" in resultado.detalhe


def test_pasta_processados_nao_e_reconsumida(tmp_path):
    cfg = config_de_teste(tmp_path)
    webhook_dir = Path(cfg.whatsapp.webhook_dir)
    (webhook_dir / "processados").mkdir(parents=True, exist_ok=True)
    ja_processado = webhook_dir / "processados" / "antigo.json"
    ja_processado.write_text(json.dumps(envelope_whatsapp()), encoding="utf-8")

    resultado = canais.ler_webhook_whatsapp(cfg)

    assert resultado.arquivos == ()
    assert ja_processado.is_file()


def test_detalhe_do_whatsapp_nunca_contem_o_token(tmp_path):
    cfg = config_de_teste(tmp_path)
    webhook_dir = Path(cfg.whatsapp.webhook_dir)
    webhook_dir.mkdir(parents=True, exist_ok=True)
    (webhook_dir / "msg.json").write_text(json.dumps(envelope_whatsapp()), encoding="utf-8")
    resultado = canais.ler_webhook_whatsapp(cfg)
    assert WHATSAPP_TOKEN not in resultado.detalhe
    assert WHATSAPP_TOKEN not in resultado.arquivos[0].name
    assert WHATSAPP_TOKEN not in str(resultado.destino)


# --------------------------------------------------------------------- coletar()


def test_coletar_devolve_um_resultado_por_canal_e_ignora_desconhecido(tmp_path):
    cfg = config_de_teste(tmp_path)
    resultados = canais.coletar(cfg, canais=("telegram", "fax"), transporte=TransporteFalso({"ok": True, "result": []}))
    por_canal = {r.canal: r for r in resultados}
    assert set(por_canal) == {"telegram", "fax"}
    assert por_canal["fax"].arquivos == ()
    assert "desconhecido" in por_canal["fax"].detalhe


def test_coletar_sem_canais_usa_a_configuracao(tmp_path):
    cfg = config_de_teste(tmp_path, CANAIS_ATIVOS="whatsapp")
    webhook_dir = Path(cfg.whatsapp.webhook_dir)
    webhook_dir.mkdir(parents=True, exist_ok=True)
    (webhook_dir / "msg.json").write_text(json.dumps(envelope_whatsapp()), encoding="utf-8")
    resultados = canais.coletar(cfg)
    assert [r.canal for r in resultados] == ["whatsapp"]
    assert resultados[0].mensagens == 1


def test_coletar_whatsapp_usa_a_inbox_do_canal(tmp_path):
    cfg = config_de_teste(tmp_path, CANAIS_ATIVOS="whatsapp", INBOX_DIR=str(tmp_path / "inbox_custom"))
    webhook_dir = Path(cfg.whatsapp.webhook_dir)
    webhook_dir.mkdir(parents=True, exist_ok=True)
    (webhook_dir / "msg.json").write_text(json.dumps(envelope_whatsapp()), encoding="utf-8")
    resultado = canais.coletar(cfg)[0]
    assert resultado.arquivos[0].parent == tmp_path / "inbox_custom" / "whatsapp"


# --------------------------------------------------------------------- receptor de webhook (F8)


def _python_do_projeto() -> str:
    return str(RAIZ_PROJETO / ".venv" / "Scripts" / "python.exe")


def _http(metodo: str, porta: int, caminho: str, corpo=None):
    import urllib.error
    import urllib.request

    url = f"http://127.0.0.1:{porta}{caminho}"
    dados = None if corpo is None else json.dumps(corpo, ensure_ascii=False).encode("utf-8")
    requisicao = urllib.request.Request(url, data=dados, method=metodo)
    if dados:
        requisicao.add_header("Content-Type", "application/json")
    try:
        with urllib.request.urlopen(requisicao, timeout=15) as resposta:
            return resposta.status, resposta.read().decode("utf-8", errors="replace")
    except urllib.error.HTTPError as exc:
        return exc.code, exc.read().decode("utf-8", errors="replace")


def test_receptor_de_webhook_local_ciclo_completo(tmp_path):
    """Sobe o receptor de verdade em 127.0.0.1, faz GET/POST e alimenta o pipeline."""
    caminho_env = env_de_teste(tmp_path)
    porta = porta_livre()
    processo = subprocess.Popen(
        [
            _python_do_projeto(),
            str(RAIZ_PROJETO / "tools" / "receber_webhook_whatsapp.py"),
            "--host",
            "127.0.0.1",
            "--porta",
            str(porta),
            "--env",
            str(caminho_env),
            "--uma-vez",
        ],
        cwd=str(tmp_path),
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    try:
        esperar_no_ar(porta)

        status_ok, corpo_ok = _http(
            "GET", porta, f"/webhook/whatsapp?hub.mode=subscribe&hub.verify_token={WHATSAPP_VERIFY}&hub.challenge=DESAFIO123"
        )
        status_ruim, _ = _http(
            "GET", porta, "/webhook/whatsapp?hub.mode=subscribe&hub.verify_token=TOKEN_ERRADO&hub.challenge=X"
        )
        status_post, _ = _http("POST", porta, "/webhook/whatsapp", envelope_whatsapp())

        saida, _ = processo.communicate(timeout=30)
    finally:
        if processo.poll() is None:  # pragma: no cover - so se travar
            processo.kill()
            processo.communicate(timeout=10)

    assert status_ok == 200 and corpo_ok == "DESAFIO123"
    assert status_ruim == 403, "token de verificacao errado tem de ser recusado"
    assert status_post == 200
    assert processo.returncode == 0

    gravados = sorted(Path(tmp_path / "webhook").glob("*.json"))
    assert len(gravados) == 1, f"o receptor nao gravou o envelope: {gravados}"
    assert json.loads(gravados[0].read_text(encoding="utf-8")) == envelope_whatsapp()

    assert WHATSAPP_VERIFY not in saida, "o verify token apareceu na saida do receptor"
    assert WHATSAPP_TOKEN not in saida
    assert "verify_token: configurado" in saida
    assert "403 hub.verify_token nao confere" in saida

    # ... e a coleta leva esse envelope para o inbox, no formato do mock
    cfg = CFG.carregar(env_path=caminho_env, ambiente={})
    resultado = canais.ler_webhook_whatsapp(cfg)
    assert resultado.mensagens == 1
    coletado = ler_jsonl(resultado.arquivos[0])[0]
    mock = ler_jsonl(MOCKS / "whatsapp" / "whatsapp_1.jsonl")[0]
    assert set(coletado) == set(mock)
    assert coletado["entry"][0]["changes"][0]["value"]["messages"][0]["id"] == "wamid.COLETADO"


# --------------------------------------------------------------------- modo real: ponta a ponta


def test_modo_real_coleta_e_roda_o_pipeline_com_stub_local(tmp_path):
    """O teste que prova "o modo real liga pela configuracao", sem credencial real.

    Encadeamento exercitado de verdade: `app.run --real` -> `canais.coletar` -> urllib
    contra o stub local -> envelope no inbox -> pipeline -> planilha/fila de excecoes.
    """
    inbox = tmp_path / "inbox"
    shutil.copytree(MOCKS / "pdf", inbox / "pdf")  # NF nativa real do corpus sintetico
    (inbox / "whatsapp").mkdir(parents=True)
    (inbox / "telegram").mkdir(parents=True)

    resposta = {
        "ok": True,
        "result": [
            update_telegram(1, chat_id=TELEGRAM_CHAT_ID, texto="Pedido 7501 confirmado, total R$ 480,00"),
            update_telegram(2, chat_id="-999999999", texto="Pedido 9999 total R$ 1,00 (outra conversa)"),
        ],
    }
    with stub_http(lambda caminho, consulta, corpo: (200, resposta)) as (porta, recebidas):
        caminho_env = env_de_teste(
            tmp_path,
            TELEGRAM_API_BASE=f"http://127.0.0.1:{porta}",
            CANAIS_ATIVOS="telegram",
            INBOX_DIR=str(inbox),
        )
        resultado = subprocess.run(
            [
                _python_do_projeto(),
                "-m",
                "app.run",
                "--real",
                "--env",
                str(caminho_env),
                "--inbox",
                str(inbox),
                "--out",
                str(tmp_path / "out"),
                "--db",
                str(tmp_path / "out" / "pipeline.db"),
            ],
            cwd=str(tmp_path),
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            env={**os.environ, "PYTHONPATH": str(RAIZ_PROJETO)},
            timeout=300,
        )
        assert recebidas, "o modo real nao chegou a chamar o canal"
        assert recebidas[0]["caminho"] == f"/bot{TELEGRAM_TOKEN}/getUpdates"

    assert resultado.returncode == 0, f"STDOUT:\n{resultado.stdout}\nSTDERR:\n{resultado.stderr}"
    assert "Modo     : real" in resultado.stdout
    assert "Coleta   : telegram: 1 mensagem(ns)" in resultado.stdout, resultado.stdout

    coletados = sorted((inbox / "telegram").glob("telegram_coleta_*.jsonl"))
    assert len(coletados) == 1
    assert len(ler_jsonl(coletados[0])) == 1, "so o chat autorizado pode entrar no inbox"

    saida_completa = resultado.stdout + resultado.stderr
    assert TELEGRAM_TOKEN not in saida_completa, "token vazou na saida do modo real"
    assert WHATSAPP_TOKEN not in saida_completa

    # trilha de auditoria: a NF do corpus foi publicada, a ordem coletada foi para revisao
    registros = [json.loads(linha) for linha in (tmp_path / "out" / "auditoria.jsonl").read_text(encoding="utf-8").splitlines() if linha.strip()]
    do_telegram = [r for r in registros if "telegram_coleta" in str(r["artefato"])]
    assert len(do_telegram) == 1, "o envelope coletado nao aparece na trilha de auditoria"
    assert do_telegram[0]["numero_pedido"] == "7501", "a mensagem coletada nao virou artefato do pedido"
    assert do_telegram[0]["canal"] == "telegram"
    assert do_telegram[0]["acao"] == "revisao", (
        "ordem vinda de mensagem nao entra na planilha (sem itens o total nao reconcilia): "
        "tem de ir para revisao humana"
    )

    da_nota = [r for r in registros if str(r["artefato"]).endswith("FORN-ALFA_nf_1001.pdf")]
    assert da_nota and da_nota[0]["acao"] == "inserido"

    from openpyxl import load_workbook

    planilha = load_workbook(tmp_path / "out" / "controle_financeiro.xlsx").active
    cabecalho = [celula.value for celula in planilha[1]]
    linhas = [
        dict(zip(cabecalho, [celula.value for celula in planilha[n]]))
        for n in range(2, planilha.max_row + 1)
        if any(celula.value is not None for celula in planilha[n])
    ]
    numeros = {str(linha["numero_pedido"]) for linha in linhas}
    assert "1001" in numeros, "a NF do corpus nao foi processada pelo modo real"
    assert "7501" not in numeros

    fila = json.loads((tmp_path / "out" / "fila_excecoes.json").read_text(encoding="utf-8"))
    assert fila["total"] >= 1, "a ordem coletada tem de estar na fila de excecoes"

    log = sorted((tmp_path / "logs").glob("pipeline-*.log"))
    assert log, "o modo real tem de escrever o log em arquivo"
    texto_log = log[0].read_text(encoding="utf-8")
    assert "rodada iniciada" in texto_log and "modo=real" in texto_log
    assert TELEGRAM_TOKEN not in texto_log, "token vazou no arquivo de log"
